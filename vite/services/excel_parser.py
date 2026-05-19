"""Lectura de archivos Excel de entrada y escritura del informe de salida."""

from pathlib import Path
import re
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

from vite.config.settings import FORMATOS_SOPORTADOS, FORMATO_REGEX, CAMPOS_TERCEROS


def read_excel(path: Path) -> pd.DataFrame:
    """
    Lee un archivo Excel (.xls o .xlsx) y retorna un DataFrame con strings puros.

    - Usa engine='xlrd' para .xls y engine='openpyxl' para .xlsx.
    - Lee todas las celdas como str (dtype=str) para preservar ceros a la izquierda.
    - Convierte celdas vacías / NaN a None.
    - Hace strip de espacios en columnas de texto.

    Args:
        path: ruta al archivo Excel.
    Returns:
        pd.DataFrame con los datos del archivo.
    """
    suffix = path.suffix.lower()
    if suffix == ".xls":
        engine = "xlrd"
    else:
        engine = "openpyxl"

    df = pd.read_excel(path, dtype=str, engine=engine)

    # Reemplazar NaN / cadenas vacías resultantes del cast por None
    df = df.where(pd.notna(df), None)

    # Strip de espacios en columnas de tipo object
    df = df.apply(lambda col: col.str.strip() if col.dtype == object else col)

    return df


def extract_formato_from_filename(filename: str) -> int | None:
    """
    Extrae el número de formato DIAN a partir del nombre de archivo.

    Busca el primer número que coincida con FORMATO_REGEX dentro del nombre.

    Args:
        filename: nombre del archivo (con o sin extensión).
    Returns:
        int del formato encontrado, o None si no hubo coincidencia.
    """
    match = re.search(FORMATO_REGEX, filename)
    if match:
        return int(match.group())
    return None


def write_report(
    output_path: Path,
    hoja1_data: list[dict],
    hoja2_data: list[dict] | None,
) -> None:
    """
    Escribe el informe de validaciones en un archivo Excel con dos hojas.

    Hoja 1 - "Informe de Validaciones":
        Columnas: NMDOC | TPDOC | Tipo Acción | Código Validación | Descripción
        Color de fila según tipo_accion:
            ACTINF → amarillo (#FFFF00)
            VALRES → rojo claro (#FFAAAA)
            INFOOK → verde claro (#AAFFAA)

    Hoja 2 - "Datos Validados" (solo si hoja2_data no es None):
        Columnas: TPDOC | NMDOC | DV | AP1 | AP2 | NM1 | NM2 | RZ | DIR | DPTO | MPIO | PAIS
        Encabezados con fondo azul claro (#AAD4FF).

    Args:
        output_path: ruta de destino del archivo Excel generado.
        hoja1_data: lista de dicts con claves nmdoc, tpdoc, tipo_accion,
                    codigo_validacion, descripcion.
        hoja2_data: lista de dicts con claves lowercase de CAMPOS_TERCEROS,
                    o None para omitir la hoja.
    """
    wb = Workbook()

    # ------------------------------------------------------------------ #
    # Hoja 1 — Informe de Validaciones
    # ------------------------------------------------------------------ #
    ws1 = wb.active
    ws1.title = "Informe de Validaciones"

    # Estilos de relleno
    COLOR_HEADER = PatternFill("solid", fgColor="D0D0D0")
    COLOR_ACTINF = PatternFill("solid", fgColor="FFFF00")
    COLOR_VALRES = PatternFill("solid", fgColor="FFAAAA")
    COLOR_INFOOK = PatternFill("solid", fgColor="AAFFAA")
    COLOR_MAP = {
        "ACTINF": COLOR_ACTINF,
        "VALRES": COLOR_VALRES,
        "INFOOK": COLOR_INFOOK,
    }

    headers1 = ["NMDOC", "TPDOC", "Tipo Acción", "Código Validación", "Descripción"]
    for col_idx, h in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col_idx, value=h)
        cell.font = Font(bold=True)
        cell.fill = COLOR_HEADER

    for row_idx, item in enumerate(hoja1_data, 2):
        ws1.cell(row=row_idx, column=1, value=item.get("nmdoc"))
        ws1.cell(row=row_idx, column=2, value=item.get("tpdoc"))
        ws1.cell(row=row_idx, column=3, value=item.get("tipo_accion"))
        ws1.cell(row=row_idx, column=4, value=item.get("codigo_validacion"))
        ws1.cell(row=row_idx, column=5, value=item.get("descripcion"))
        color = COLOR_MAP.get(item.get("tipo_accion", ""), None)
        if color:
            for col_idx in range(1, 6):
                ws1.cell(row=row_idx, column=col_idx).fill = color

    # Ajuste automático de ancho de columnas
    for col in ws1.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws1.column_dimensions[get_column_letter(col[0].column)].width = min(
            max_len + 2, 50
        )

    # ------------------------------------------------------------------ #
    # Hoja 2 — Datos Validados (condicional)
    # ------------------------------------------------------------------ #
    if hoja2_data is not None:
        ws2 = wb.create_sheet("Datos Validados")
        COLOR_HEADER2 = PatternFill("solid", fgColor="AAD4FF")

        # CAMPOS_TERCEROS ya viene en el orden correcto desde settings
        headers2 = CAMPOS_TERCEROS  # ["TPDOC", "NMDOC", "DV", ...]

        for col_idx, h in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col_idx, value=h)
            cell.font = Font(bold=True)
            cell.fill = COLOR_HEADER2

        for row_idx, tercero in enumerate(hoja2_data, 2):
            for col_idx, campo in enumerate(headers2, 1):
                ws2.cell(
                    row=row_idx,
                    column=col_idx,
                    value=tercero.get(campo.lower()),
                )

        # Ajuste automático de ancho de columnas (máximo 30)
        for col in ws2.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws2.column_dimensions[get_column_letter(col[0].column)].width = min(
                max_len + 2, 30
            )

    wb.save(output_path)
